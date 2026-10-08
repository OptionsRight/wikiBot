"""Deterministic synthetic business data, never company documents."""
from pathlib import Path
import sys
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw


def make(root):
    root.mkdir(parents=True, exist_ok=True)
    image = Image.new('RGB', (600, 150), 'white')
    ImageDraw.Draw(image).text((20, 40), 'Scanned limit: 80 USD', fill='black')
    image.save(root / 'scan.png')
    for boundary in (False, True):
        name = 'boundary' if boundary else 'normal'
        doc = Document()
        doc.add_heading('Synthetic reimbursement policy', 0)
        doc.add_paragraph('Travel reimbursement limit is 80 USD per day.')
        if boundary:
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged heading: USD'
            table.cell(1, 0).text = 'Travel'
            table.cell(1, 1).text = '80'
            doc.add_picture(str(root / 'scan.png'), width=Inches(3))
        doc.save(root / f'{name}.docx')
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[1])
        slide.shapes.title.text = 'Travel policy'
        slide.placeholders[1].text = 'Travel reimbursement limit is 80 USD per day.'
        if boundary:
            slide.shapes.add_picture(str(root / 'scan.png'), Inches(1), Inches(4), width=Inches(3))
            ppt.slides.add_slide(ppt.slide_layouts[6])
        ppt.save(root / f'{name}.pptx')
        book = Workbook()
        sheet = book.active
        sheet.title = 'Limits'
        sheet.append(['Category', 'USD per day'])
        sheet.append(['Travel', 80])
        if boundary:
            sheet['B3'] = '=B2*2'
            sheet.merge_cells('A4:B4')
            sheet['A4'] = 'Merged footnote'
            hidden = book.create_sheet('Hidden')
            hidden.sheet_state = 'hidden'
            hidden['A1'] = 'Not approved for publication'
        book.save(root / f'{name}.xlsx')
        pdf = canvas.Canvas(str(root / f'{name}.pdf'), invariant=1)
        pdf.drawString(60, 760, 'Travel reimbursement limit is 80 USD per day.')
        if boundary:
            pdf.showPage()
            pdf.drawImage(str(root / 'scan.png'), 60, 600, width=450, height=112)
        pdf.save()
    for ext in ('pdf', 'docx', 'pptx', 'xlsx'):
        (root / f'invalid.{ext}').write_bytes(b'not a document')


if __name__ == '__main__':
    make(Path(sys.argv[1]))
